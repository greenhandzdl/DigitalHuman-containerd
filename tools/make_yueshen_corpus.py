"""生成 yueshen_rag 容器的语料（构建期跑，产物只进镜像层，不落进仓库）。

yueshen_rag 的 CorpusLoader 只认 .pdf/.docx（server.py:195），而 `fay` 仓库里一份都没有
（`find fay -name '*.pdf' -o -name '*.docx'` 排掉 .venv 是 0 个），上游 README 说的默认语料
目录 `新知识库` 也不存在。所以「知识库能不能真入库真检索」这件事缺的不是代码，是一份语料。
这里就地生成一份 —— 用同一个 python-docx 写，正好也验了依赖层里的 docx 解析路径。

MARKER 这个字符串是 `probes/yueshen_probe.py` 里断言要检索回来的那个词，两份文件各写
一遍、没有共享常量：**改这里必须同时改探针**，否则那条判据会直接判红
（检索不到标记词 = 向量库里的东西和探针以为在的东西不是同一份），不会假绿。
"""

import sys
from pathlib import Path

from docx import Document

MARKER = "悦肾探针语料YS20260417"

# 每段一个可回答的事实，标记词嵌在第二段里 —— 检索回来的必须正是那一段。
PARAGRAPHS = [
    "随访规范总则：本语料由 containerd 在构建 yueshen-rag 镜像时生成，仅用于验证"
    "「文档解析 → 切块 → 向量化 → 入库 → 检索」这五步在容器里真的跑通。",
    f"编号 {MARKER} 的专项随访要求：每次随访结束后，必须在 48 小时内把血压与体重录回系统，"
    "超过 48 小时未录视为一次漏访，需要由责任护士在次日 10 点前补录并说明原因。",
    "通用提醒：老人用药提醒的默认提前量是 15 分钟，训练计划提醒的默认提前量是 30 分钟。",
]


def main() -> int:
    out_dir = Path(sys.argv[1] if len(sys.argv) > 1 else "/app/corpus")
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("康养随访知识库（容器化验证语料）", level=1)
    for text in PARAGRAPHS:
        doc.add_paragraph(text)
    target = out_dir / "followup_probe_corpus.docx"
    doc.save(str(target))
    print(f"[corpus] {target} {target.stat().st_size} 字节，{len(PARAGRAPHS)} 段，标记词 {MARKER!r}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
